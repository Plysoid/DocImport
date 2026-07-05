from pathlib import Path
from openpyxl import Workbook, load_workbook


class CorrectionExporter:
    """Створює діагностичні файли для наповнення ТоварКор.xlsx / АдресиКор.xlsx."""

    def ensure_templates(self, data_dir):
        data_dir = Path(data_dir)
        data_dir.mkdir(parents=True, exist_ok=True)
        self._ensure_product_template(data_dir / "ТоварКор.xlsx")
        self._ensure_address_template(data_dir / "АдресиКор.xlsx")

    def export_unknowns(self, invoices, output_dir):
        output_dir = Path(output_dir)
        output_dir.mkdir(parents=True, exist_ok=True)
        self._unknown_products(invoices, output_dir / "unknown_products.xlsx")
        self._unknown_addresses(invoices, output_dir / "unknown_addresses.xlsx")
        self._warnings(invoices, output_dir / "warnings.xlsx")
        self._fixes(invoices, output_dir / "fixes.xlsx")

    def _ensure_product_template(self, filename):
        if filename.exists():
            return
        wb = Workbook()
        ws = wb.active
        ws.title = "ТоварКор"
        ws.append(["ITEM", "ITEMNAMENAKL"])
        wb.save(filename)

    def _ensure_address_template(self, filename):
        if filename.exists():
            return
        wb = Workbook()
        ws = wb.active
        ws.title = "АдресиКор"
        ws.append(["Код", "АдресаНакл"])
        wb.save(filename)

    def _unknown_products(self, invoices, filename):
        rows = {}
        for inv in invoices:
            for it in inv.items:
                if it.product is None:
                    key = (it.raw_name or "", it.barcode or "")
                    rows[key] = [
                        "", it.raw_name or "", it.barcode or "",
                        inv.doc or "", inv.source_file or "", inv.page
                    ]
        wb = Workbook()
        ws = wb.active
        ws.title = "unknown_products"
        ws.append(["ITEM", "ITEMNAMENAKL", "ITEMSHK_NAKL", "DOC", "FILE", "PAGE"])
        for row in rows.values():
            ws.append(row)
        wb.save(filename)

    def _unknown_addresses(self, invoices, filename):
        rows = {}
        for inv in invoices:
            if inv.shop is None:
                key = inv.address or ""
                rows[key] = ["", inv.address or "", inv.doc or "", inv.source_file or "", inv.page]
        wb = Workbook()
        ws = wb.active
        ws.title = "unknown_addresses"
        ws.append(["Код", "АдресаНакл", "DOC", "FILE", "PAGE"])
        for row in rows.values():
            ws.append(row)
        wb.save(filename)

    def _warnings(self, invoices, filename):
        wb = Workbook()
        ws = wb.active
        ws.title = "warnings"
        ws.append(["DOC", "FILE", "PAGE", "SHOP", "ADDRESS", "DOCQTY", "ITEMQTY", "QTYDIFF", "DOCSUM", "ITEMSUM", "SUMDIFF", "MATCH", "NOTE"])
        for inv in invoices:
            item_sum = round(sum(float(it.amount or 0) for it in inv.items), 2)
            item_qty = round(sum(float(it.qty or 0) for it in inv.items), 3)
            doc_sum = round(float(getattr(inv, "doc_sum", 0) or 0), 2)
            doc_qty = round(float(getattr(inv, "doc_qty", 0) or 0), 3)
            sum_diff = round(item_sum - doc_sum, 2) if doc_sum else 0
            qty_diff = round(item_qty - doc_qty, 3) if doc_qty else 0
            methods = "; ".join(sorted(set(str(it.match_method or "") for it in inv.items if it.match_method)))
            has_warn = any(x in methods for x in ("docsum", "docqty", "amountfix-docsum", "qtyfix-docqty"))
            if has_warn or (doc_sum and abs(sum_diff) > 0.10) or (doc_qty and abs(qty_diff) > 0.001):
                ws.append([
                    inv.doc or "",
                    inv.source_file or "",
                    inv.page,
                    inv.shop.code if inv.shop else "",
                    inv.address or "",
                    doc_qty,
                    item_qty,
                    qty_diff,
                    doc_sum,
                    item_sum,
                    sum_diff,
                    methods,
                    "Перевірка кількості та суми накладної",
                ])
        wb.save(filename)

    def _fixes(self, invoices, filename):
        wb = Workbook()
        ws = wb.active
        ws.title = "fixes"
        ws.append(["DOC", "FILE", "PAGE", "ITEM", "ITEMNAME", "ITEMNAKL", "FIELD", "OLD", "NEW", "REASON"])
        for inv in invoices:
            for fx in getattr(inv, "fixes", []):
                ws.append([
                    fx.get("doc", ""),
                    fx.get("file", ""),
                    fx.get("page", ""),
                    fx.get("item", ""),
                    fx.get("item_name", ""),
                    fx.get("item_nakl", ""),
                    fx.get("field", ""),
                    fx.get("old", ""),
                    fx.get("new", ""),
                    fx.get("reason", ""),
                ])
        wb.save(filename)
