from pathlib import Path
import re
from openpyxl import load_workbook
from rapidfuzz import fuzz, process
from models.entities import Shop
from utils.text import norm


DEBUG_ADDRESS = True
ADDRESS_DEBUG_LOG = Path("logs") / "address_debug.txt"


BAD_LINE_MARKERS = (
    "КОКА", "БЕВЕР", "IBAN", "IВАМ", "ІВАМ", "ЇВАМ", "ЄДРПОУ",
    "ПОСТАВКА", "НАКЛАДНА", "ЗАВАНТАЖЕН", "ЦЕНТР", "ВОДІЙ",
    "УВАГА", "ТЕЛЕФОНУЙТЕ", "РЕКВІЗИТ", "СІТІБАНК", "СИТIБАНК",
    "АКТ РОЗБІЖ", "ПІДПИС", "ВАНТАЖ ОТРИМАВ", "ПОДАТКОВА",
    "ВСЬОГО", "ПДВ", "ТАРА", "ВІДПУСК", "ДОРУЧЕННЯ",
)

HEADER_STOP_MARKERS = (
    "ТОВАРНИЙ РОЗДІЛ", "ШТРИХ", "УВАГА", "ЦІНИ ВКАЗАНО",
)

POSTAL_RE = re.compile(r"\b\d{5}\s*,\s*[А-ЯІЇЄҐA-Z'’\- ]+")
POSTAL_START_RE = re.compile(r"\b\d{5}\s*,")
PHONE_PREFIX_RE = re.compile(r"^.*?(?:ТЕЛ\.?|TEN)\s*[: ]*\d{6,}\s*", re.I)
HOUSE_RE = re.compile(r"\b\d{1,4}(?:\s*[А-ЯA-Z])?(?:/\d{1,4})?\b")


class ShopCatalog:
    def __init__(self):
        self.items: list[Shop] = []
        self.by_code: dict[str, Shop] = {}
        self.by_norm: dict[str, Shop] = {}
        self._keys: list[tuple[str, Shop, list[str]]] = []
        self.corr_addresses: dict[str, Shop] = {}

    def load(self, filename):
        p = Path(filename)
        if not p.exists():
            raise FileNotFoundError(f"Не знайдено довідник адрес: {p}")

        wb = load_workbook(p, read_only=True, data_only=True)
        ws = wb.active
        rows = ws.iter_rows(values_only=True)
        header = [str(x or "").strip().upper() for x in next(rows)]

        def try_col(*names):
            for n in names:
                if n.upper() in header:
                    return header.index(n.upper())
            return None

        c_code = try_col("SHOP", "КОД", "КОД МАГАЗИНУ")
        c_addr = try_col("ADDRESS", "АДРЕСА", "АДРЕС")
        if c_code is None:
            c_code = 0
        if c_addr is None:
            c_addr = 1

        for r in rows:
            code = str(r[c_code] or "").strip()
            addr = str(r[c_addr] or "").strip()
            if not code or not addr:
                continue
            sh = Shop(code=code.zfill(3)[-3:], address=addr)
            self.items.append(sh)
            self.by_code[sh.code] = sh
            key = norm(addr)
            self.by_norm[key] = sh
            self._keys.append((key, sh, self._tokens(key)))
        wb.close()
        if DEBUG_ADDRESS:
            try:
                ADDRESS_DEBUG_LOG.parent.mkdir(parents=True, exist_ok=True)
                ADDRESS_DEBUG_LOG.write_text("", encoding="utf-8")
            except Exception:
                pass
        self._load_corrections(p.parent / "АдресиКор.xlsx")
        self._debug("LOADED SHOPS:", len(self.items), "CORRECTIONS:", len(self.corr_addresses))

    def _load_corrections(self, filename: Path):
        if not filename.exists():
            return
        wb = load_workbook(filename, read_only=True, data_only=True)
        ws = wb.active
        rows = ws.iter_rows(values_only=True)
        try:
            first = next(rows)
            header = [str(x or "").strip().upper() for x in first]
        except StopIteration:
            wb.close(); return

        def idx(*names):
            for n in names:
                if n.upper() in header:
                    return header.index(n.upper())
            return None

        c_code = idx("КОД", "SHOP", "КОД МАГАЗИНУ")
        c_addr = idx("АДРЕСАНАКЛ", "АДРЕСА НАКЛ", "ADDRESSNAKL", "АДРЕСА")

        # Підтримка АдресиКор.xlsx без шапки: код у першій колонці, адреса у другій.
        data_rows = rows
        if c_code is None or c_addr is None:
            c_code, c_addr = 0, 1
            data_rows = iter([first] + list(rows))

        for r in data_rows:
            if len(r) <= max(c_code, c_addr):
                continue
            code = str(r[c_code] or "").strip()
            addr = str(r[c_addr] or "").strip()
            if not code or not addr:
                continue
            sh = self._shop_by_code(code)
            if sh:
                # АдресиКор може містити кілька варіантів через |:
                # повний рядок з індексом і короткий рядок без міста.
                # Зберігаємо всі варіанти, щоб збіг спрацював і по
                # "81652, НОВИЙ РОЗДІЛ ...", і по "ТАРАСА ШЕВЧЕНКА 16".
                variants = [addr]
                variants.extend(str(addr).split("|"))
                for v in variants:
                    v = self._clean_address_segment(v)
                    k = norm(v)
                    # Не додаємо в корекції службове OCR-сміття типу "І".
                    # Інакше майже будь-яка адреса, що містить літеру І,
                    # може помилково збігтися з цією корекцією.
                    if not k:
                        continue
                    if len(k) < 8:
                        continue
                    if not self._address_numbers(k):
                        continue
                    self.corr_addresses[k] = sh
        wb.close()

    def _shop_by_code(self, code):
        c = str(code).strip().zfill(3)[-3:]
        return self.by_code.get(c) or Shop(code=c, address="")

    def _tokens(self, s: str):
        skip = {
            "М", "МІСТО", "ЛЬВIВ", "ЛЬВІВ", "ЛЬВОВ", "ВУЛ", "ВУЛИЦЯ",
            "ПР", "ПР.", "ПРОСП", "ПРОСПЕКТ", "БУЛ", "БУЛЬВАР", "И", "І", "I",
            "ТАРАСА", "СТЕПАНА",
            "79000", "79005", "79040", "79052", "79060", "79066",
        }
        return [x for x in norm(s).split() if len(x) >= 2 and x not in skip and not re.fullmatch(r"\d{5,}", x)]

    def _address_numbers(self, s: str):
        """Повертає номери будинків без індексів і телефонів.

        5+ цифр не вважаємо номером будинку. Також виправляємо
        типовий OCR-випадок у номері: 5ЗА -> 53А.
        """
        text = norm(s)
        # OCR: кирилична З у номері будинку часто означає цифру 3.
        text = re.sub(r"(?<=\d)З(?=[А-ЯA-Z]?\b)", "3", text)

        nums = []
        for x in re.findall(r"\d+[А-ЯA-Z]?(?:/\d+)?", text):
            if re.fullmatch(r"\d{5,}", x):
                continue
            nums.append(x)
        return set(nums)

    def _clean_address_segment(self, s: str) -> str:
        s = str(s or "").strip()
        s = PHONE_PREFIX_RE.sub("", s).strip()
        # Інколи OCR вставляє службовий код 07442 після адреси.
        s = re.split(r"\b07442\b", s, maxsplit=1)[0]
        # Прибираємо службові хвости після адреси.
        s = re.split(r"\b(?:АКТ|ПІДПИС|ВАНТАЖ|ЄДРПОУ|IBAN|ІВАМ|ЇВАМ|IВАМ|БАНК|ВІДПУСК|ВСЬОГО|ПДВ)\b", s, maxsplit=1, flags=re.I)[0]
        s = s.strip(" |;:,.©")
        s = re.sub(r"\s+", " ", s)
        return s

    def _is_bad_line(self, line: str) -> bool:
        u = str(line or "").upper()
        return any(b in u for b in BAD_LINE_MARKERS)

    def _header_lines(self, text: str):
        out = []
        for raw in str(text or "").splitlines():
            line = raw.strip()
            if not line:
                continue
            u = line.upper()
            if any(m in u for m in HEADER_STOP_MARKERS):
                break
            out.append(line)
        return out

    def _line_address_candidates(self, line: str):
        out = []
        if not line:
            return out

        # Якщо в рядку є один або кілька індексів, беремо фрагменти після кожного індексу.
        # Це прибирає адресу постачальника на кшталт '... ГОРОДНИЦЬКА 43, 79019 79052, ЛЬВІВ ...'.
        starts = [m.start() for m in POSTAL_START_RE.finditer(line)]
        if starts:
            for st in starts:
                out.append(line[st:])

        # Також розбиваємо по вертикальних OCR-розділювачах і беремо короткі адресні фрагменти.
        for part in re.split(r"\|", line):
            part = self._clean_address_segment(part)
            if not part or self._is_bad_line(part):
                continue
            if HOUSE_RE.search(part):
                out.append(part)

        clean = self._clean_address_segment(line)
        if clean and not self._is_bad_line(clean) and HOUSE_RE.search(clean):
            out.append(clean)

        return out

    def _candidate_addresses(self, text: str):
        """Витягує короткі кандидати адреси з шапки без ручних прив'язок до конкретних вулиць."""
        lines = self._header_lines(text)
        candidates = []

        def add(s):
            s = self._clean_address_segment(s)
            if not s or self._is_bad_line(s):
                return
            if len(s) > 180:
                return
            candidates.append(s)

        for i, line in enumerate(lines):
            # Кандидати з самого рядка.
            for c in self._line_address_candidates(line):
                add(c)

            # Вікна 2-3 рядки: місто/індекс часто в одному рядку, вулиця в наступному.
            for window in (2, 3):
                if i + window <= len(lines):
                    block_lines = []
                    ok = False
                    for j in range(i, i + window):
                        cl = self._clean_address_segment(lines[j])
                        if not cl or self._is_bad_line(cl):
                            continue
                        block_lines.append(cl)
                        if POSTAL_START_RE.search(cl) or HOUSE_RE.search(cl):
                            ok = True
                    if ok and block_lines:
                        add(" ".join(block_lines))

        # Додатково: якщо парсер уже передав коротку адресу, залишаємо її кандидатом.
        if text and len(str(text)) < 220:
            s = self._clean_address_segment(str(text))
            if s and HOUSE_RE.search(s):
                candidates.append(s)
        # Унікальні по norm, зберігаючи порядок.
        seen = set()
        out = []
        for c in candidates:
            k = norm(c)
            if k and k not in seen:
                seen.add(k)
                out.append(c)
        return out

    def _best_corr_address(self, candidates):
        if not self.corr_addresses:
            return None
        cand_norms = [norm(c) for c in candidates if norm(c)]
        for cn in cand_norms:
            cand_nums = self._address_numbers(cn)
            for key, sh in self.corr_addresses.items():
                if not key:
                    continue
                key_nums = self._address_numbers(key)
                # Для корекцій номер будинку має збігатися точно.
                # Інакше В. Великого 51 може помилково попасти в корекцію 67.
                if self._house_number_ok(cand_nums, key_nums) is not True:
                    continue
                n = len(key)
                if cn == key or cn.startswith(key) or key.startswith(cn) or key in cn or cn in key:
                    return sh
                # Порівнюємо фрагменти такої ж довжини, як АдресаНакл.
                # Для коротких кандидатів також порівнюємо token_set_ratio:
                # "ТАРАСА ШЕВЧЕНКА 16" має збігатися з короткою корекцією.
                if fuzz.token_set_ratio(cn, key) >= 92:
                    return sh
                for i in range(0, max(1, len(cn) - n + 1)):
                    frag = cn[i:i+n]
                    if frag == key or fuzz.ratio(frag, key) >= 92:
                        return sh
        return None

    def _house_number_ok(self, cand_nums, shop_nums):
        if not shop_nums:
            return True
        if not cand_nums:
            return False
        for cn in cand_nums:
            for sn in shop_nums:
                cm = re.match(r"\d+", cn)
                sm = re.match(r"\d+", sn)
                if cm and sm and cm.group(0) == sm.group(0):
                    return True
        for cn in cand_nums:
            for sn in shop_nums:
                cm = re.match(r"\d+", cn)
                sm = re.match(r"\d+", sn)
                if cm and sm:
                    cb, sb = cm.group(0), sm.group(0)
                    # OCR може губити першу цифру: 17 -> 7. Це слабкий збіг.
                    if len(cb) == 1 and len(sb) > 1 and sb.endswith(cb):
                        return "weak"
        return False

    def _score(self, candidate: str, shop_key: str, shop_tokens: list[str]):
        nt = norm(candidate)
        if not nt or not shop_key:
            return 0

        nums_shop = self._address_numbers(shop_key)
        nums_cand = self._address_numbers(nt)
        num_ok = self._house_number_ok(nums_cand, nums_shop)

        cand_tokens = set(self._tokens(nt))
        shop_token_set = set(shop_tokens)
        if not shop_token_set or not cand_tokens:
            return 0

        common = cand_tokens & shop_token_set
        coverage = len(common) / len(shop_token_set)

        # Номер будинку — ключовий запобіжник від помилкових адрес.
        if num_ok is False:
            return 0

        if shop_key in nt or nt in shop_key:
            return 100 if num_ok is True else 88

        if shop_tokens and all(t in nt for t in shop_tokens):
            return 98 if num_ok is True else 87

        if num_ok is True:
            if coverage >= 0.80:
                return 92
            if coverage >= 0.60:
                return 86
        elif num_ok == "weak" and coverage >= 0.65:
            return 87

        ratio = fuzz.token_set_ratio(nt, shop_key)
        if ratio >= 92 and (num_ok is True or num_ok == "weak"):
            return ratio if num_ok is True else min(ratio, 88)
        return 0

    def _debug(self, *parts):
        if not DEBUG_ADDRESS:
            return
        try:
            ADDRESS_DEBUG_LOG.parent.mkdir(parents=True, exist_ok=True)
            with ADDRESS_DEBUG_LOG.open("a", encoding="utf-8") as f:
                f.write(" ".join(str(x) for x in parts) + "\n")
        except Exception:
            pass

    def _debug_address_state(self, text, candidates, corr=None, best=None):
        if not DEBUG_ADDRESS:
            return
        self._debug("\n" + "=" * 80)
        self._debug("ADDR INPUT:", str(text or "")[:500].replace("\n", " | "))
        self._debug("ADDR CANDIDATES:")
        for c in candidates:
            self._debug("  -", c, "=>", norm(c), "nums=", sorted(self._address_numbers(c)))
        self._debug("ADDR CORR KEYS COUNT:", len(self.corr_addresses))
        keys497 = [k for k, v in self.corr_addresses.items() if getattr(v, "code", "") == "497"]
        self._debug("HAS 497 CORR:", keys497)
        if corr:
            self._debug("CORR MATCH:", getattr(corr, "code", ""), getattr(corr, "address", ""))
        if best and best[0] is not None:
            self._debug("BEST FUZZY:", getattr(best[0], "code", ""), getattr(best[0], "address", ""), best[1], best[2])
        elif best:
            self._debug("BEST FUZZY:", best)

    def find(self, text: str):
        candidates = self._candidate_addresses(text)

        corr = self._best_corr_address(candidates)
        if corr:
            self._debug_address_state(text, candidates, corr=corr)
            return corr, "addr-corr"

        best = (None, "not-found", 0)
        for cand in candidates:
            for key, sh, toks in self._keys:
                score = self._score(cand, key, toks)
                if score > best[2]:
                    best = (sh, f"addr:{score:.0f}", score)

        self._debug_address_state(text, candidates, best=best)
        if best[0] is not None and best[2] >= 86:
            return best[0], best[1]
        return None, "not-found"
