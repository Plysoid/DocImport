from pathlib import Path
from openpyxl import load_workbook
from rapidfuzz import process, fuzz
from models.entities import Product
from utils.text import norm


class ProductCatalog:
    def __init__(self):
        self.items: list[Product] = []
        self.by_code = {}
        self.by_barcode = {}
        self.by_barcode12 = {}
        self.names = {}
        self.corr_names = {}  # norm(ITEMNAMENAKL) -> Product

    def load(self, filename):
        p = Path(filename)
        if not p.exists():
            raise FileNotFoundError(f"Не знайдено довідник товарів: {p}")
        wb = load_workbook(p, read_only=True, data_only=True)
        ws = wb.active
        rows = ws.iter_rows(values_only=True)
        header = [str(x or "").strip().upper() for x in next(rows)]

        def col(*names):
            for n in names:
                if n.upper() in header:
                    return header.index(n.upper())
            raise Exception(f"У Товари.xlsx немає колонки з варіантів: {names}")

        c_code = col("ITEM", "КОД")
        c_name = col("ITEMNAME", "ТОВАР")
        c_bar = col("ITEMSHK", "ШТРИХ-КОД", "ШТРИХКОД")
        for r in rows:
            code = str(r[c_code] or "").strip()
            name = str(r[c_name] or "").strip()
            bar_raw = str(r[c_bar] or "").strip().replace(".0", "")
            bar = ''.join(ch for ch in bar_raw if ch.isdigit())

            # В основному довіднику приймаємо тільки товари з коректним EAN-8 або EAN-13.
            # Інші записи повністю ігноруємо, щоб не було хибних збігів по назві.
            if len(bar) not in (8, 13):
                continue
            if not code or not name:
                continue

            pr = Product(code=code, name=name, barcode=bar)
            self.items.append(pr)
            self.by_code[code] = pr
            self.by_barcode[bar] = pr
            if len(bar) == 13:
                self.by_barcode12[bar[:12]] = pr
            self.names[norm(name)] = pr
        wb.close()
        self._load_corrections(p.parent / "ТоварКор.xlsx")

    def _load_corrections(self, filename: Path):
        if not filename.exists():
            return
        wb = load_workbook(filename, read_only=True, data_only=True)
        ws = wb.active
        rows = ws.iter_rows(values_only=True)
        try:
            header = [str(x or "").strip().upper() for x in next(rows)]
        except StopIteration:
            wb.close(); return

        def idx(*names):
            for n in names:
                if n.upper() in header:
                    return header.index(n.upper())
            return None

        c_item = idx("ITEM", "КОД")
        c_nakl = idx("ITEMNAMENAKL", "НАЗВАНАКЛ", "НАЗВА НАКЛ")
        if c_item is None or c_nakl is None:
            wb.close(); return

        for r in rows:
            code = str(r[c_item] or "").strip()
            nakl = str(r[c_nakl] or "").strip()
            if not code or not nakl:
                continue
            pr = self.by_code.get(code)
            if pr:
                self.corr_names[norm(nakl)] = pr
        wb.close()

    def _best_corr_name(self, name: str):
        n = norm(name)
        if not n or not self.corr_names:
            return None
        if n in self.corr_names:
            return self.corr_names[n]
        for key, pr in self.corr_names.items():
            if len(key) >= 6 and (key in n or n in key):
                return pr
        best = process.extractOne(n, list(self.corr_names.keys()), scorer=fuzz.WRatio)
        if best and best[1] >= 90:
            return self.corr_names[best[0]]
        return None

    def _best_barcode_distance(self, b: str):
        b = ''.join(ch for ch in str(b or "") if ch.isdigit())
        if len(b) < 12:
            return None
        key = b[:12]
        hits = []
        for k12, pr in self.by_barcode12.items():
            if len(k12) != 12:
                continue
            dist = sum(1 for a, c in zip(key, k12) if a != c)
            if dist <= 1:
                hits.append(pr)
        uniq = {}
        for pr in hits:
            uniq[pr.code] = pr
        if len(uniq) == 1:
            return next(iter(uniq.values()))
        return None

    def find(self, barcode: str, name: str):
        # 0. Примусове співставлення назви з накладної: data\ТоварКор.xlsx
        pr = self._best_corr_name(name)
        if pr:
            return pr, "corr-name"

        b = ''.join(ch for ch in str(barcode or "") if ch.isdigit())

        if b and b in self.by_barcode:
            return self.by_barcode[b], "barcode"

        if len(b) >= 12:
            key12 = b[:12]
            if key12 in self.by_barcode12:
                return self.by_barcode12[key12], "barcode-12"

        if len(b) == 12 and b in self.by_barcode12:
            return self.by_barcode12[b], "barcode-12"

        pr = self._best_barcode_distance(b)
        if pr:
            return pr, "barcode-12-hamming1"

        if b and len(b) >= 7:
            for k, pr in self.by_barcode.items():
                kd = ''.join(ch for ch in k if ch.isdigit())
                if len(kd) <= 8 and (b == kd or b in kd or kd in b):
                    return pr, "barcode-short"

        n = norm(name)
        if n in self.names:
            return self.names[n], "name"
        if n and self.names:
            best = process.extractOne(n, list(self.names.keys()), scorer=fuzz.WRatio)
            if best and best[1] >= 72:
                return self.names[best[0]], f"fuzzy:{best[1]:.0f}"
        return None, "not-found"
