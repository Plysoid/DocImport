from pathlib import Path

from openpyxl import Workbook, load_workbook


class PriceHistory:
    """Остання підтверджена ціна товару.

    Файл має дві колонки:
    - ITEM: код товару з довідника;
    - LASTPRICE: остання ціна, яка підтвердилась хоча б у двох різних накладних.
    """

    def __init__(self, path):
        self.path = Path(path)
        self.prices = {}

    def _item_code(self, value):
        """Нормалізує ITEM із Excel до текстового коду довідника.

        Excel може віддати 7688 як int, float або текст. Для історії це має бути
        один і той самий ключ: "7688".
        """
        if value is None:
            return ""
        if isinstance(value, float) and value.is_integer():
            return str(int(value))
        return str(value).strip()

    def _price(self, value):
        """Читає LASTPRICE і з числа Excel, і з тексту виду 385,32."""
        if value is None:
            return 0.0
        try:
            return round(float(value), 2)
        except Exception:
            pass

        s = str(value).strip().replace(" ", "")
        if not s:
            return 0.0

        # Український формат: 1.234,56 або 385,32
        if "," in s:
            s = s.replace(".", "").replace(",", ".")
        try:
            return round(float(s), 2)
        except Exception:
            return 0.0

    def load(self):
        self.prices = {}
        if not self.path.exists():
            return

        wb = load_workbook(self.path, data_only=True)
        ws = wb.active
        headers = [str(c.value or "").strip().upper() for c in ws[1]]
        try:
            item_col = headers.index("ITEM") + 1
            price_col = headers.index("LASTPRICE") + 1
        except ValueError:
            return

        for row in range(2, ws.max_row + 1):
            item = self._item_code(ws.cell(row, item_col).value)
            if not item:
                continue
            price = self._price(ws.cell(row, price_col).value)
            if price > 0:
                self.prices[item] = price

    def get(self, item_code):
        return self.prices.get(self._item_code(item_code))

    def update(self, confirmed_prices):
        changed = False
        for item, price in confirmed_prices.items():
            item = self._item_code(item)
            price = self._price(price)
            if not item or price <= 0:
                continue
            if self.prices.get(item) != price:
                self.prices[item] = price
                changed = True
        if changed:
            self.save()

    def save(self):
        self.path.parent.mkdir(parents=True, exist_ok=True)
        wb = Workbook()
        ws = wb.active
        ws.title = "price_history"
        ws.append(["ITEM", "LASTPRICE"])
        for item in sorted(self.prices):
            ws.append([item, self.prices[item]])
        wb.save(self.path)
